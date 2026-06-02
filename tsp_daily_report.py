"""Daily Excel reporting for current-driver TSP routes."""

from __future__ import annotations

from copy import deepcopy
from datetime import date, datetime
import json
import os
import threading
from typing import Any, Dict, Iterable, List, Optional

from config import MainConfig


_HISTORY_LOCK = threading.Lock()


def _clean_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _parse_float(value: Any, default: float = 0.0) -> float:
    try:
        if value in (None, ""):
            return default
        return float(value)
    except (TypeError, ValueError):
        return default


def _parse_int(value: Any, default: int = 0) -> int:
    try:
        if value in (None, ""):
            return default
        return int(float(value))
    except (TypeError, ValueError):
        return default


def _today_string() -> str:
    return date.today().isoformat()


def tsp_history_file(config: MainConfig) -> str:
    api_config = getattr(config, "api", None)
    configured = _clean_text(getattr(api_config, "tsp_daily_report_history_file", ""))
    if configured:
        return os.path.abspath(configured)
    return os.path.abspath(os.path.join(os.getcwd(), "logs", "tsp_routes_history.jsonl"))


def tsp_report_output_dir(config: MainConfig) -> str:
    api_config = getattr(config, "api", None)
    configured = _clean_text(getattr(api_config, "tsp_daily_report_output_dir", ""))
    if configured:
        return os.path.abspath(configured)

    output_config = getattr(config, "output", None)
    excel_dir = _clean_text(getattr(output_config, "excel_output_dir", ""))
    if excel_dir:
        return os.path.abspath(excel_dir)
    return os.path.abspath(os.path.join(os.getcwd(), "output", "excel"))


def tsp_daily_report_path(config: MainConfig, report_date: Optional[str] = None) -> str:
    day = report_date or _today_string()
    return os.path.join(tsp_report_output_dir(config), f"tsp_daily_report_{day}.xlsx")


def record_tsp_result(config: MainConfig, result: Dict[str, Any], payload: Optional[Dict[str, Any]] = None) -> None:
    """Append one successful /tsp result to the JSONL history."""
    if not isinstance(result, dict) or result.get("status") != "ok":
        return

    now = datetime.now()
    record = {
        "recorded_at": now.isoformat(timespec="seconds"),
        "service_date": now.date().isoformat(),
        "driver_id": _clean_text(result.get("driver_id")),
        "driver_name": _clean_text(result.get("driver_name")),
        "metric": _clean_text(result.get("metric")),
        "customers_count": _parse_int(result.get("customers_count")),
        "total_distance_km": _parse_float(result.get("total_distance_km")),
        "total_time_minutes": _parse_float(result.get("total_time_minutes")),
        "total_quantity": _parse_float(result.get("total_quantity")),
        "total_turnover": _parse_float(result.get("total_turnover")),
        "start_location": deepcopy(result.get("start_location")),
        "end_location": deepcopy(result.get("end_location")),
        "map_file": _clean_text(result.get("map_file")),
        "map_upload": deepcopy(result.get("map_upload")),
        "tsp_process": deepcopy(result.get("tsp_process")),
        "skipped_customers": deepcopy(result.get("skipped_customers") or []),
        "delivery_order": deepcopy(result.get("delivery_order") or []),
    }

    if isinstance(payload, dict):
        record["request_driver_id"] = _clean_text(payload.get("driver_id"))

    path = tsp_history_file(config)
    os.makedirs(os.path.dirname(path), exist_ok=True)
    line = json.dumps(record, ensure_ascii=False)
    with _HISTORY_LOCK:
        with open(path, "a", encoding="utf-8") as fh:
            fh.write(line + "\n")


def load_tsp_records(config: MainConfig, report_date: Optional[str] = None) -> List[Dict[str, Any]]:
    day = report_date or _today_string()
    path = tsp_history_file(config)
    if not os.path.exists(path):
        return []

    records: List[Dict[str, Any]] = []
    with _HISTORY_LOCK:
        with open(path, "r", encoding="utf-8") as fh:
            for line in fh:
                line = line.strip()
                if not line:
                    continue
                try:
                    record = json.loads(line)
                except json.JSONDecodeError:
                    continue
                if record.get("service_date") == day:
                    records.append(record)
    return records


def _format_duration(minutes: Any) -> str:
    total = int(round(_parse_float(minutes)))
    return f"{total // 60:02d}:{total % 60:02d}"


def _coords_text(value: Any) -> str:
    if isinstance(value, dict):
        lat = value.get("lat")
        lon = value.get("lon")
        if lat is not None and lon is not None:
            return f"{lat},{lon}"
    if isinstance(value, (list, tuple)) and len(value) >= 2:
        return f"{value[0]},{value[1]}"
    return ""


def _upload_status_text(value: Any) -> str:
    if not isinstance(value, dict):
        return ""
    if value.get("attempted") is False:
        return "не е опитано"
    if "saved_count" in value:
        return f"качени {value.get('saved_count')}"
    if "status_code" in value:
        return f"HTTP {value.get('status_code')}"
    if value.get("error"):
        return f"грешка: {value.get('error')}"
    return "опитано"


def _auto_width(ws) -> None:
    for column_cells in ws.columns:
        max_len = 0
        column_letter = column_cells[0].column_letter
        for cell in column_cells:
            text = str(cell.value or "")
            max_len = max(max_len, min(len(text), 80))
        ws.column_dimensions[column_letter].width = max(10, min(max_len + 2, 55))


def _style_header(ws) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill

    fill = PatternFill("solid", fgColor="4472C4")
    font = Font(bold=True, color="FFFFFF")
    for cell in ws[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"


def generate_tsp_daily_excel_report(config: MainConfig, report_date: Optional[str] = None) -> str:
    """Generate an Excel report for all TSP routes recorded on report_date."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    day = report_date or _today_string()
    records = load_tsp_records(config, day)
    output_path = tsp_daily_report_path(config, day)
    os.makedirs(os.path.dirname(output_path), exist_ok=True)

    wb = Workbook()
    ws_summary = wb.active
    ws_summary.title = "Обобщение"

    total_routes = len(records)
    total_customers = sum(_parse_int(record.get("customers_count")) for record in records)
    total_distance = sum(_parse_float(record.get("total_distance_km")) for record in records)
    total_time = sum(_parse_float(record.get("total_time_minutes")) for record in records)
    total_quantity = sum(_parse_float(record.get("total_quantity")) for record in records)
    total_turnover = sum(_parse_float(record.get("total_turnover")) for record in records)

    summary_rows = [
        ("Дата", day),
        ("Брой TSP маршрути", total_routes),
        ("Общо клиенти", total_customers),
        ("Общо количество", round(total_quantity, 2)),
        ("Общо оборот", round(total_turnover, 2)),
        ("Общо км", round(total_distance, 3)),
        ("Общо време", _format_duration(total_time)),
        ("Генериран на", datetime.now().isoformat(timespec="seconds")),
    ]
    for row in summary_rows:
        ws_summary.append(row)
    ws_summary["A1"].font = Font(bold=True)
    ws_summary["B1"].font = Font(bold=True)
    _auto_width(ws_summary)

    ws_routes = wb.create_sheet("Маршрути")
    ws_routes.append(
        [
            "Дата/час",
            "ID шофьор",
            "Име",
            "Клиенти",
            "Количество",
            "Оборот",
            "Км",
            "Време",
            "Време (мин)",
            "Старт",
            "Край",
            "Карта",
            "Upload",
            "Процес",
        ]
    )
    for record in records:
        ws_routes.append(
            [
                record.get("recorded_at", ""),
                record.get("driver_id", ""),
                record.get("driver_name", ""),
                _parse_int(record.get("customers_count")),
                _parse_float(record.get("total_quantity")),
                _parse_float(record.get("total_turnover")),
                _parse_float(record.get("total_distance_km")),
                _format_duration(record.get("total_time_minutes")),
                _parse_float(record.get("total_time_minutes")),
                _coords_text(record.get("start_location")),
                _coords_text(record.get("end_location")),
                record.get("map_file", ""),
                _upload_status_text(record.get("map_upload")),
                (record.get("tsp_process") or {}).get("mode", ""),
            ]
        )
    _style_header(ws_routes)
    _auto_width(ws_routes)

    api_config = getattr(config, "api", None)
    if bool(getattr(api_config, "tsp_daily_report_include_details", True)):
        ws_details = wb.create_sheet("Клиенти")
        ws_details.append(
            [
                "Дата/час",
                "ID шофьор",
                "Име шофьор",
                "Ред",
                "ID клиент",
                "Клиент",
                "Документ",
                "GPS",
                "Количество",
                "Оборот",
                "Пристигане",
                "Тръгване",
                "Работно време",
                "Статус прозорец",
                "Км от предишен",
                "Км общо",
                "Път (мин)",
                "Чакане (мин)",
                "Коментар",
            ]
        )
        for record in records:
            for stop in record.get("delivery_order") or []:
                coords = _coords_text(stop.get("coordinates"))
                ws_details.append(
                    [
                        record.get("recorded_at", ""),
                        record.get("driver_id", ""),
                        record.get("driver_name", ""),
                        _parse_int(stop.get("sequence")),
                        stop.get("customer_id", ""),
                        stop.get("customer_name", ""),
                        stop.get("order", ""),
                        coords,
                        _parse_float(stop.get("quantity")),
                        _parse_float(stop.get("turnover")),
                        stop.get("arrival_time", ""),
                        stop.get("departure_time", ""),
                        stop.get("work_time", ""),
                        stop.get("time_window_status", ""),
                        _parse_float(stop.get("distance_from_previous_km")),
                        _parse_float(stop.get("cumulative_distance_km")),
                        _parse_float(stop.get("travel_time_minutes")),
                        _parse_float(stop.get("wait_minutes")),
                        stop.get("comment", ""),
                    ]
                )
        _style_header(ws_details)
        for row in ws_details.iter_rows(min_row=2):
            row[-1].alignment = Alignment(wrap_text=True, vertical="top")
        _auto_width(ws_details)

    ws_skipped = wb.create_sheet("Пропуснати")
    ws_skipped.append(["Дата/час", "ID шофьор", "Причина", "Данни"])
    for record in records:
        for skipped in record.get("skipped_customers") or []:
            reason = skipped.get("reason", "") if isinstance(skipped, dict) else ""
            ws_skipped.append(
                [
                    record.get("recorded_at", ""),
                    record.get("driver_id", ""),
                    reason,
                    json.dumps(skipped, ensure_ascii=False) if isinstance(skipped, dict) else str(skipped),
                ]
            )
    _style_header(ws_skipped)
    _auto_width(ws_skipped)

    wb.save(output_path)
    return output_path


def report_already_exists(config: MainConfig, report_date: Optional[str] = None) -> bool:
    return os.path.exists(tsp_daily_report_path(config, report_date))
