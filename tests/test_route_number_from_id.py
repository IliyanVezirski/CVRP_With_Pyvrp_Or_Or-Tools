"""Public route labels follow the ID suffix, not the solver's list position."""
from types import SimpleNamespace
from unittest.mock import patch

import pytest
from openpyxl import Workbook

from config import MainConfig, OutputConfig
from input_handler import Customer
from output_handler import ExcelExporter
from setdata_client import build_set_data_rows
from vehicle_numbering import format_route_bus_number, route_output_number
from tests.test_output_map_visibility import VehicleMapGroupingTests


def routes_fixture():
    routes = []
    for number in (13, 3, 4):
        client = Customer(str(number), 'Client', (42.71, 23.3), 2, '42.71,23.3', document=f'000{number}')
        route = VehicleMapGroupingTests._route(2 if number == 4 else 1, f'course-{number}', [client])
        route.vehicle_key = 'van-3' if number in (3, 4) else 'van-13'
        route.schedule_id = f'10045010{number:02d}'
        routes.append(route)
    return routes


@pytest.mark.parametrize('number', [3, 13, 20, 101])
def test_reserved_number_is_suffix_not_ordinal(number):
    route = SimpleNamespace(schedule_id=f'10045010{number:02d}')
    assert route_output_number(OutputConfig(), 0, route, [route]) == number


def test_automatic_center_and_unused_reservations():
    cfg = OutputConfig(excel_bus_number_start=3, center_bus_numbering_enabled=True,
                       center_bus_numbering_start_id='1004501015')
    normal = SimpleNamespace(vehicle_type='internal_bus', reserved_schedule_ids=('1004501003',))
    center = SimpleNamespace(vehicle_type='center_bus')
    routes = [normal, center]
    assert [route_output_number(cfg, i, r, routes) for i, r in enumerate(routes)] == [4, 15]
    assert format_route_bus_number(cfg, 0, normal, routes) == '1004501004'


def test_other_prefix_is_not_guessed_and_saturday_uses_its_prefix():
    route = SimpleNamespace(schedule_id='1004501213')
    assert route_output_number(OutputConfig(), 1, route) == 2
    cfg = OutputConfig(excel_bus_number_prefix='100450121', excel_bus_number_digits=1)
    assert route_output_number(cfg, 1, route) == 3


def test_setdata_keeps_stop_number_and_uses_id_suffix_for_bukva():
    cfg = MainConfig()
    cfg.output = OutputConfig()
    cfg.set_data.set_data_id_grafik_template = '{bus_number}'
    cfg.set_data.set_data_bukva_template = 'БХ{route_number}-{stop_number}'
    routes = routes_fixture()
    rows = build_set_data_rows(SimpleNamespace(routes=routes), cfg)
    assert [(r['IdGrafik'], r['Bukva']) for r in rows] == [
        ('1004501013', 'БХ13-1'), ('1004501003', 'БХ3-1'), ('1004501004', 'БХ4-1')]
    assert [r.trip_number for r in routes] == [1, 1, 2]


def test_report_route_columns_match_business_id():
    routes = routes_fixture()
    exporter = ExcelExporter(OutputConfig())
    cfg = VehicleMapGroupingTests._config_without_zones()
    with patch('output_handler.get_config', return_value=cfg):
        wb = Workbook()
        exporter._create_routes_sheet(wb, SimpleNamespace(routes=routes))
        exporter._create_vehicle_stats_sheet(wb, SimpleNamespace(routes=routes))
    for ws in list(wb.worksheets)[1:]:
        headers = {c.value: c.column for c in ws[1]}
        assert [ws.cell(row=i, column=headers['Маршрут']).value for i in (2, 3, 4)] == [13, 3, 4]
        assert [ws.cell(row=i, column=headers['ID бус']).value for i in (2, 3, 4)] == [
            '1004501013', '1004501003', '1004501004']


@pytest.mark.parametrize('provider', ['google', 'osm'])
def test_maps_keep_id_route_and_course_consistent(provider):
    routes = routes_fixture()
    generator = VehicleMapGroupingTests._map_generator(provider)
    cfg = VehicleMapGroupingTests._config_without_zones()
    with patch('output_handler.get_config', return_value=cfg):
        payload = generator._build_google_routes(routes, all_routes=routes)
        assert [r['routeNumber'] for r in payload] == [13, 3, 4]
        assert [r['outputNumber'] for r in payload] == ['1004501013', '1004501003', '1004501004']
        assert [r['tripNumber'] for r in payload] == [1, 1, 2]
        document = generator.create_single_route_map(routes[0], 1, (42.7, 23.3), routes)
        rendered = document.get_root().render() if hasattr(document, 'get_root') else document.html_content
        assert 'маршрут 13' in rendered.lower() or '"singleRouteNumber": 13' in rendered
        comment = generator._route_delivery_order_html_comment(routes[0], 1, routes)
        assert 'route_number=13' in comment and 'bus_id=1004501013' in comment


@pytest.mark.parametrize('provider', ['google', 'osm'])
def test_combined_map_does_not_renumber_second_course_or_automatic_id(provider):
    routes = routes_fixture()
    generator = VehicleMapGroupingTests._map_generator(provider)
    generator.config.excel_bus_number_start = 5
    # Make the first route automatic; it must remain ...05, not become ...09
    # when a visible number passes through nested map-generation functions.
    routes[0].schedule_id = ''
    cfg = VehicleMapGroupingTests._config_without_zones()
    with patch('output_handler.get_config', return_value=cfg):
        payload = generator._build_google_routes(routes, all_routes=routes)
        assert payload[0]['routeNumber'] == 5
        assert payload[0]['outputNumber'] == '1004501005'
        document = generator.create_single_route_map(routes[0], 1, (42.7, 23.3), routes)
        rendered = document.get_root().render() if hasattr(document, 'get_root') else document.html_content
        assert 'bus_id=1004501005' in document._cvrp_top_html_comment
        assert 'маршрут 5' in rendered.lower() or '"singleRouteNumber": 5' in rendered
        combined = generator.create_vehicle_route_map(routes[1:], [2, 3], (42.7, 23.3), routes)
        rendered = combined.get_root().render() if hasattr(combined, 'get_root') else combined.html_content
        if provider == 'google':
            assert '"routeNumber": 3' in rendered and '"routeNumber": 4' in rendered
        else:
            assert 'Маршрут 3' in rendered and 'Маршрут 4' in rendered
