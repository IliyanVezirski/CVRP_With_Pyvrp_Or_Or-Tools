import csv
import os
import tempfile
import unittest
from types import SimpleNamespace
from unittest.mock import Mock, patch

from openpyxl import Workbook

from input_handler import Customer
from output_handler import (
    ExcelExporter,
    OutputHandler,
    _duplicate_route_map_bus_ids,
    _route_map_file_name,
    _route_start_time_minutes,
    _stored_route_schedule_entries,
    _total_elapsed_route_minutes,
)
from setdata_client import build_set_data_rows
from vehicle_numbering import (
    count_physical_vehicles,
    format_route_bus_number,
    order_routes_for_output,
)


def _route(key=None, trip=1, route_id=None, vehicle_type="internal_bus", **overrides):
    values = {
        "vehicle_type": SimpleNamespace(value=vehicle_type),
        "vehicle_id": 0,
        "vehicle_name": "Internal Bus 1",
        "customers": [],
        "depot_location": (42.0, 23.0),
        "end_location": (42.0, 23.0),
        "total_distance_km": 0.0,
        "total_time_minutes": 0.0,
        "total_volume": 0.0,
    }
    if key is not None:
        values["vehicle_key"] = key
        values["trip_number"] = trip
        values["route_id"] = route_id or f"{key}-trip-{trip}"
    values.update(overrides)
    return SimpleNamespace(**values)


class MultiTripVehicleNumberingTests(unittest.TestCase):
    def setUp(self):
        self.output = SimpleNamespace(
            excel_bus_number_prefix="10045010",
            excel_bus_number_digits=2,
            center_bus_numbering_enabled=True,
            center_bus_numbering_start_id=1004501015,
        )

    def test_courses_are_grouped_but_receive_unique_business_numbers(self):
        a2 = _route("A", 2)
        c1 = _route("C", 1, vehicle_type="center_bus")
        b1 = _route("B", 1)
        a1 = _route("A", 1)
        c2 = _route("C", 2, vehicle_type="center_bus")

        routes = order_routes_for_output([a2, c1, b1, a1, c2])

        self.assertEqual(routes, [a1, a2, b1, c1, c2])
        self.assertEqual(count_physical_vehicles(routes), 3)
        numbers = [
            format_route_bus_number(self.output, idx, route, routes)
            for idx, route in enumerate(routes)
        ]
        self.assertEqual(numbers, ["1004501001", "1004501002", "1004501003", "1004501015", "1004501016"])

    def test_non_center_courses_skip_the_occupied_center_number_range(self):
        regular = [_route(f"R{index}", 1) for index in range(1, 17)]
        center = [
            _route("C", 1, vehicle_type="center_bus"),
            _route("C", 2, vehicle_type="center_bus"),
        ]
        routes = order_routes_for_output(regular + center)

        numbers = [
            format_route_bus_number(self.output, idx, route, routes)
            for idx, route in enumerate(routes)
        ]

        self.assertEqual(len(numbers), len(set(numbers)))
        self.assertEqual(numbers[13:18], [
            "1004501014",
            "1004501017",
            "1004501018",
            "1004501015",
            "1004501016",
        ])

    def test_multitrip_driver_id_does_not_force_courses_to_share_one_number(self):
        first = _route("A", 1, driver_id="1004501099")
        second = _route("A", 2, driver_id="1004501099")
        routes = order_routes_for_output([second, first])

        self.assertEqual(
            [format_route_bus_number(self.output, i, route, routes) for i, route in enumerate(routes)],
            ["1004501001", "1004501002"],
        )

        standalone = _route(driver_id="1004501099")
        self.assertEqual(
            format_route_bus_number(self.output, 0, standalone, [standalone]),
            "1004501099",
        )

    def test_legacy_routes_keep_the_previous_order_and_numbering(self):
        first = _route()
        center = _route(vehicle_type="center_bus", vehicle_name="Center Bus 1")
        second = _route(vehicle_name="Internal Bus 2")
        routes = order_routes_for_output([first, center, second])

        self.assertEqual(routes, [first, second, center])
        self.assertEqual(
            [format_route_bus_number(self.output, i, route, routes) for i, route in enumerate(routes)],
            ["1004501001", "1004501002", "1004501015"],
        )


class MultiTripSetDataTests(unittest.TestCase):
    def test_each_course_uses_its_unique_bus_number_as_id_grafik(self):
        first_customer = Customer(
            id="client-1",
            name="First",
            coordinates=(42.1, 23.1),
            volume=1,
            original_gps_data="42.1,23.1",
            document="doc-1",
            plas_doc="plas-1",
        )
        second_customer = Customer(
            id="client-2",
            name="Second",
            coordinates=(42.2, 23.2),
            volume=1,
            original_gps_data="42.2,23.2",
            document="doc-2",
            plas_doc="plas-2",
        )
        solution = SimpleNamespace(routes=[
            _route("physical-1", 1, customers=[first_customer]),
            _route("physical-1", 2, customers=[second_customer]),
        ])
        config = SimpleNamespace(
            set_data=SimpleNamespace(
                set_data_command="setData",
                set_data_done_flag="D",
                set_data_id_skld="128",
                set_data_id_grafik="",
                set_data_id_grafik_template="{bus_number}",
                set_data_bukva_template="{bus_number}",
                set_data_depot_id_skld_map="",
            ),
            output=SimpleNamespace(
                excel_bus_number_prefix="10045010",
                excel_bus_number_digits=2,
                center_bus_numbering_enabled=False,
            ),
            locations=None,
        )

        rows = build_set_data_rows(solution, config)

        self.assertEqual(
            [(row["IdGrafik"], row["Bukva"]) for row in rows],
            [
                ("1004501001", "1004501001"),
                ("1004501002", "1004501002"),
            ],
        )

    def test_new_route_fields_are_template_placeholders_not_http_parameters(self):
        customer = Customer(
            id="client-1",
            name="Client",
            coordinates=(42.1, 23.1),
            volume=2.5,
            original_gps_data="42.1,23.1",
            document="doc-1",
            plas_doc="plas-1",
        )
        route = _route(
            "physical-7",
            2,
            route_id="route-7-2",
            customers=[customer],
            planned_start_minutes=615,
            planned_end_minutes=735,
        )
        solution = SimpleNamespace(routes=[route])
        set_data = SimpleNamespace(
            set_data_command="setData",
            set_data_done_flag="D",
            set_data_id_skld="128",
            set_data_id_grafik="G",
            set_data_id_grafik_template="",
            set_data_bukva_template=(
                "{vehicle_key}|{trip_number}|{route_id}|{planned_start}|{planned_end}|{bus_number}"
            ),
            set_data_depot_id_skld_map="",
        )
        config = SimpleNamespace(set_data=set_data, output=SimpleNamespace(
            excel_bus_number_prefix="10045010",
            excel_bus_number_digits=2,
            center_bus_numbering_enabled=False,
        ), locations=None)

        rows = build_set_data_rows(solution, config)

        self.assertEqual(len(rows), 1)
        self.assertEqual(
            set(rows[0]),
            {"cmd", "IdPlasDoc", "DoneFlag", "IdSkld", "Bukva", "IdGrafik"},
        )
        self.assertEqual(
            rows[0]["Bukva"],
            "physical-7|2|route-7-2|615|735|1004501001",
        )


class MultiTripOutputTests(unittest.TestCase):
    def test_excel_multitrip_sheets_keep_only_the_course_column(self):
        routes = [
            _route("Маршрут#1", 1, route_id="Маршрут#1:trip:1"),
            _route("Маршрут#1", 2, route_id="Маршрут#1:trip:2"),
        ]
        solution = SimpleNamespace(
            routes=routes,
            total_distance_km=0.0,
            total_time_minutes=0.0,
        )
        workbook = Workbook()
        workbook.remove(workbook.active)
        exporter = ExcelExporter(SimpleNamespace())

        exporter._create_routes_sheet(workbook, solution)
        exporter._create_vehicle_stats_sheet(workbook, solution)

        for sheet_name in ("Маршрути", "Статистики по бусове"):
            sheet = workbook[sheet_name]
            headers = [cell.value for cell in sheet[1]]
            self.assertIn("Курс", headers)
            self.assertNotIn("ID курс", headers)
            self.assertNotIn("Ключ бус", headers)
            rendered_values = "\n".join(
                str(cell.value)
                for row in sheet.iter_rows()
                for cell in row
                if cell.value is not None
            )
            self.assertNotIn(":trip:", rendered_values)
            self.assertNotIn("Маршрут#1", rendered_values)

    def test_overview_time_includes_reload_gap_between_courses(self):
        first = _route(
            "A",
            1,
            total_time_minutes=25,
            planned_start_minutes=480,
            planned_end_minutes=505,
        )
        second = _route(
            "A",
            2,
            total_time_minutes=25,
            planned_start_minutes=535,
            planned_end_minutes=560,
        )
        other_bus = _route(
            "B",
            1,
            total_time_minutes=20,
            planned_start_minutes=480,
            planned_end_minutes=500,
        )

        # Bus A works 80 minutes including its 30-minute reload; bus B works 20.
        self.assertAlmostEqual(
            100,
            _total_elapsed_route_minutes([first, second, other_bus]),
        )

    def test_route_filenames_and_multitrip_upload_bus_ids_are_unique(self):
        first = _route("A", 1, route_id="A-1")
        second = _route("A", 2, route_id="A-2")
        routes = [first, second]
        used = set()

        first_name = _route_map_file_name(first, 1, "2026-07-11", used, routes)
        second_name = _route_map_file_name(second, 2, "2026-07-11", used, routes)
        self.assertIn("kurs_1", first_name)
        self.assertIn("kurs_2", second_name)
        self.assertNotEqual(first_name, second_name)

        output = SimpleNamespace(
            excel_bus_number_prefix="10045010",
            excel_bus_number_digits=2,
            center_bus_numbering_enabled=False,
        )
        upload_items = [
            {
                "file_path": "one.html",
                "bus_id": format_route_bus_number(output, 0, first, routes),
                "trip_number": 1,
                "route_id": "A-1",
            },
            {
                "file_path": "two.html",
                "bus_id": format_route_bus_number(output, 1, second, routes),
                "trip_number": 2,
                "route_id": "A-2",
            },
        ]
        self.assertEqual([item["bus_id"] for item in upload_items], ["1004501001", "1004501002"])
        self.assertEqual(_duplicate_route_map_bus_ids(upload_items), [])

        with tempfile.TemporaryDirectory() as temp_dir:
            # The upload identity is the business bus ID, not the path.  Even
            # reusing one generated HTML file for two courses must produce two
            # aligned files[] / pData2[] entries.
            file_path = os.path.join(temp_dir, "shared.html")
            with open(file_path, "w", encoding="utf-8") as handle:
                handle.write("<html></html>")
            real_items = [
                {**item, "file_path": file_path}
                for item in upload_items
            ]
            handler = OutputHandler(SimpleNamespace(
                route_maps_upload_mode="effect",
                route_maps_upload_url="https://effect.invalid/upload",
                route_maps_upload_token_field="pData",
                route_maps_upload_token="token",
                route_maps_upload_file_field="files[]",
                route_maps_upload_bus_id_field="pData2[]",
                route_maps_upload_timeout_seconds=10,
            ))
            response = Mock(status_code=200, text="OK")
            response.raise_for_status.return_value = None
            with patch("output_handler.requests.post", return_value=response) as request_post:
                result = handler._upload_route_maps(real_items)
            request_post.assert_called_once()
            posted_data = request_post.call_args.kwargs["data"]
            posted_files = request_post.call_args.kwargs["files"]
            self.assertIn(("pData2[]", "1004501001"), posted_data)
            self.assertIn(("pData2[]", "1004501002"), posted_data)
            self.assertEqual(2, sum(1 for field, _payload in posted_files if field == "files[]"))
        self.assertEqual(result, "HTTP 200: OK")

    def test_single_trip_upload_keeps_the_original_numeric_bus_id(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            file_path = os.path.join(temp_dir, "one.html")
            with open(file_path, "w", encoding="utf-8") as handle:
                handle.write("<html></html>")
            handler = OutputHandler(SimpleNamespace(
                route_maps_upload_mode="effect",
                route_maps_upload_url="https://effect.invalid/upload",
                route_maps_upload_token_field="pData",
                route_maps_upload_token="token",
                route_maps_upload_file_field="files[]",
                route_maps_upload_bus_id_field="pData2[]",
                route_maps_upload_timeout_seconds=10,
            ))
            response = Mock(status_code=200, text="OK")
            response.raise_for_status.return_value = None
            with patch("output_handler.requests.post", return_value=response) as request_post:
                result = handler._upload_route_maps([
                    {"file_path": file_path, "bus_id": "1004501001", "trip_number": 1}
                ])
            posted_data = request_post.call_args.kwargs["data"]
            self.assertIn(("pData2[]", "1004501001"), posted_data)
            self.assertNotIn("trip", str(posted_data).lower())
        self.assertEqual(result, "HTTP 200: OK")

    def test_planned_start_and_physical_course_totals_reach_reports(self):
        first = _route("A", 1, planned_start_minutes=480, planned_end_minutes=600)
        second = _route("A", 2, planned_start_minutes="10:30", planned_end_minutes="12:15")
        self.assertEqual(_route_start_time_minutes(first), 480)
        self.assertEqual(_route_start_time_minutes(second), 630)

        second.schedule_entries = [{
            "index": 1,
            "start_time_minutes": 480,
            "arrival_time_minutes": 510,
            "total_time_with_start": 525,
        }]
        shifted = _stored_route_schedule_entries(second)
        self.assertEqual(shifted[0]["start_time_minutes"], 630)
        self.assertEqual(shifted[0]["arrival_time_minutes"], 660)
        self.assertEqual(shifted[0]["total_time_with_start"], 675)

        solution = SimpleNamespace(
            routes=[first, second],
            total_distance_km=0.0,
            total_time_minutes=0.0,
        )
        exporter = ExcelExporter(SimpleNamespace())
        workbook = Workbook()
        workbook.remove(workbook.active)
        exporter._create_summary_sheet(workbook, solution, [])
        sheet = workbook["Обобщение"]
        stats = {
            str(sheet.cell(row=row, column=1).value): sheet.cell(row=row, column=2).value
            for row in range(1, sheet.max_row + 1)
        }
        self.assertEqual(stats["Брой физически бусове"], 1)
        self.assertEqual(stats["Брой курсове"], 2)

    def test_multitrip_csv_has_course_identity_and_second_course_start(self):
        customer = Customer("c2", "Second", (42.1, 23.1), 1.0, "42.1,23.1")
        first = _route("A", 1, planned_start_minutes=480)
        second = _route("A", 2, route_id="A-second", planned_start_minutes=630, customers=[customer])
        solution = SimpleNamespace(routes=[first, second])

        with tempfile.TemporaryDirectory() as temp_dir:
            file_path = os.path.join(temp_dir, "routes.csv")
            exporter = ExcelExporter(SimpleNamespace(csv_output_file=file_path))
            result = exporter.export_routes_csv(solution)
            with open(result, newline="", encoding="utf-8-sig") as handle:
                rows = list(csv.DictReader(handle))

        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["ID бус"], "1004501002")
        self.assertEqual(rows[0]["Курс"], "2")
        self.assertEqual(rows[0]["ID курс"], "A-second")
        self.assertEqual(rows[0]["Ключ бус"], "A")
        self.assertEqual(rows[0]["Стартово време (мин)"], "630")


if __name__ == "__main__":
    unittest.main()
