from copy import deepcopy
from types import SimpleNamespace
from unittest.mock import patch

import pytest
from openpyxl import load_workbook

import config_gui
import cvrp_api_server as api
from cvrp_solver import CVRPSolution
from input_handler import InputHandler
from missing_gps_settings import MISSING_GPS_DEFAULTS
from output_handler import ExcelExporter
from setdata_client import build_missing_gps_set_data_rows, upload_solution_set_data
from tests.test_missing_gps_unserved import config, record


def test_missing_gps_is_sent_once_after_grouping_and_never_with_1973():
    cfg = config()
    cfg.set_data.enable_make_group = True
    cfg.set_data.set_data_unserved_done_flag = "1975"
    cfg.set_data.set_data_unserved_final_done_flag = "1973"
    customers = InputHandler(cfg).load_data_from_json_records([
        record("gps-missing"), record("ordinary", GPS="42.7,23.3"),
    ]).customers
    with patch("setdata_client.send_set_data_row", return_value={"status_code": 200}) as send:
        result = upload_solution_set_data(
            SimpleNamespace(routes=[], dropped_customers=[]), cfg,
            SimpleNamespace(warehouse_customers=customers),
        )
    rows = [call.args[0] for call in send.call_args_list]
    absent = [row for row in rows if row.get("IdPlasDoc") == "plas-gps-missing"]
    assert len(absent) == 1
    assert absent[0]["DoneFlag"] == "0"
    assert absent[0]["IdGrafik"] == "NO-GPS"
    assert absent[0]["Bukva"] == "GPS-1"
    assert rows[-1] == absent[0]
    assert [r["DoneFlag"] for r in rows if r.get("IdPlasDoc") == "plas-ordinary"] == ["1975", "1973"]
    assert result["missing_gps_attempted"] == 1
    assert result["unserved_done_flag"]["attempted"] == 1


def test_missing_gps_uses_independent_flag_and_upload_toggle():
    cfg = config()
    cfg.set_data.enable_unserved_set_data_upload = False
    cfg.set_data.set_data_missing_gps_done_flag = "72"
    customers = InputHandler(cfg).load_data_from_json_records([record()]).customers
    solution = SimpleNamespace(routes=[], dropped_customers=[])
    allocation = SimpleNamespace(warehouse_customers=customers)
    with patch("setdata_client.send_set_data_row", return_value={}) as send:
        upload_solution_set_data(solution, cfg, allocation)
    assert send.call_count == 1
    assert send.call_args.args[0]["DoneFlag"] == "72"
    cfg.set_data.enable_missing_gps_set_data_upload = False
    with patch("setdata_client.send_set_data_row") as send:
        upload_solution_set_data(solution, cfg, allocation)
    send.assert_not_called()


def test_missing_gps_only_does_not_trigger_makegroup():
    cfg = config()
    cfg.set_data.enable_make_group = True
    customers = InputHandler(cfg).load_data_from_json_records([record()]).customers
    with patch("setdata_client.send_set_data_row", return_value={}) as send:
        result = upload_solution_set_data(SimpleNamespace(routes=[], dropped_customers=[]), cfg,
                                         SimpleNamespace(warehouse_customers=customers))
    assert send.call_count == 1
    assert result["make_group"]["attempted"] == 0
    assert result["unserved_done_flag"]["attempted"] == 0


def test_legacy_defaults_and_separate_web_profile_roundtrip():
    cfg = config()
    legacy = vars(cfg.set_data).copy()
    for field in MISSING_GPS_DEFAULTS:
        legacy.pop(field)
    cfg.set_data = SimpleNamespace(**legacy)
    defaults = api._web_run_settings_from_config(cfg)["set_data"]
    assert defaults["set_data_missing_gps_done_flag"] == "0"
    assert defaults["set_data_missing_gps_id_grafik"] == "UNSERVED"
    values = dict(enable_missing_gps_set_data_upload=True, set_data_missing_gps_done_flag="55",
                  set_data_missing_gps_id_grafik="00123", set_data_missing_gps_bukva_template="BG-{stop_number}")
    with patch.object(api, "get_config", return_value=cfg):
        changed, applied, ignored = api._build_web_run_config_override({"run_settings": {"set_data": values}})
    assert not ignored
    assert {k: api._web_run_settings_from_config(changed)["set_data"][k] for k in values} == values
    assert not hasattr(cfg.set_data, "set_data_missing_gps_done_flag")


def test_desktop_save_inserts_fields_in_legacy_config(tmp_path):
    gui = config_gui.ConfigGUI.__new__(config_gui.ConfigGUI)
    gui.cfg = config()
    gui.cfg.vehicles = []
    source = 'from dataclasses import dataclass\n@dataclass\nclass SetDataConfig:\n    enable_set_data_upload: bool = False\n'
    path = tmp_path / "config.py"
    path.write_text(source, encoding="utf-8")
    values = dict(enable_missing_gps_set_data_upload=False, set_data_missing_gps_done_flag="0",
                  set_data_missing_gps_id_grafik="000099", set_data_missing_gps_bukva_template="ГП-{stop_number}")
    with patch.object(config_gui, "_base_dir", str(tmp_path)):
        assert gui._apply_to_config_file({"set_data." + key: value for key, value in values.items()})
    scope = {}
    exec(compile(path.read_text(encoding="utf-8"), str(path), "exec"), scope)
    saved = scope["SetDataConfig"]()
    assert {k: getattr(saved, k) for k in values} == values


def test_web_controls_load_and_collect_all_fields():
    cfg = config()
    html = api._web_gui_html(cfg.api, "127.0.0.1", 18095)
    for element in ("setDataMissingGpsEnabled", "setDataMissingGpsGrafik",
                    "setDataMissingGpsBukva", "setDataMissingGpsDoneFlag"):
        assert f'id="{element}"' in html
        assert f'put("{element}"' in html
        assert f'("{element}")' in html
    assert "Заявки без GPS" in html


def test_desktop_controls_expose_independent_values():
    import tkinter as tk
    from tkinter import ttk
    root = tk.Tk()
    root.withdraw()
    try:
        gui = config_gui.ConfigGUI.__new__(config_gui.ConfigGUI)
        gui.cfg = config()
        gui.root = root
        gui.widgets, gui.controls, gui.control_default_states = {}, {}, {}
        with patch.object(gui, "_make_scrollable_frame", return_value=ttk.Frame(root)), \
             patch.object(gui, "_bind_text_editing"), patch.object(gui, "_bind_combobox_scrolling"):
            gui._add_set_data_tab(ttk.Notebook(root))
        for field in MISSING_GPS_DEFAULTS:
            assert gui.widgets["set_data." + field].get() == getattr(gui.cfg.set_data, field)
    finally:
        root.destroy()


@pytest.mark.parametrize("legacy_export", [False, True])
@pytest.mark.parametrize("only_missing", [False, True])
def test_excel_keeps_gps_orders_in_unserved_with_comment_and_setdata_destination(tmp_path, legacy_export, only_missing):
    cfg = config()
    cfg.output.excel_output_dir = str(tmp_path)
    cfg.set_data.set_data_missing_gps_id_grafik = "000123"
    cfg.set_data.set_data_missing_gps_done_flag = "0"
    rows = [record("00000001"), record("00000002")]
    if not only_missing:
        rows.append(record("normal", GPS="42.7,23.3"))
    customers = InputHandler(cfg).load_data_from_json_records(rows).customers
    exporter = ExcelExporter(cfg.output, cfg.set_data)
    if legacy_export:
        path = exporter.export_warehouse_orders(customers)
    else:
        empty = CVRPSolution([], [], 0, 0, 0, 0, True)
        path = exporter.export_all_to_single_excel(empty, customers)
    workbook = load_workbook(path)
    assert "Без GPS" not in workbook.sheetnames
    assert "Необслужени клиенти" in workbook.sheetnames
    rows = list(workbook["Необслужени клиенти"].values)
    all_reported = [dict(zip(rows[0], values)) for values in rows[1:]]
    assert len(all_reported) == (2 if only_missing else 3)
    reported = [r for r in all_reported if r["Коментар"] == "Без GPS"]
    expected = build_missing_gps_set_data_rows(customers[:2], cfg)
    assert [r["Номер поръчка"] for r in reported] == ["00000001", "00000002"]
    for report, sent in zip(reported, expected, strict=True):
        for label, key in (("Направление", "IdGrafik"), ("Буква", "Bukva"), ("Склад", "IdSkld"), ("DoneFlag", "DoneFlag")):
            assert report[label] == sent[key]
    workbook.close()
