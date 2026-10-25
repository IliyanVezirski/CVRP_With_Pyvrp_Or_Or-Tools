"""Display labels, legacy config compatibility and map utilisation: no solver run."""
import io
import json
import re
import tempfile
from pathlib import Path
import unittest
from copy import deepcopy
from types import SimpleNamespace
from unittest.mock import patch

from openpyxl import load_workbook
import cvrp_api_server as api
from config import OutputConfig, get_config
from display_units import volume_unit_label, validate_volume_unit_label, total_volume_unit_label, customer_volume_unit_label
from output_handler import _route_load_text
from tests import test_excel_report as excel_tests
from tests.test_excel_report import customer, route
from tests import test_output_map_visibility as map_tests


class DisplayUnitTests(unittest.TestCase):
    def test_independent_contexts_share_defaults_and_keep_custom_text_exact(self):
        cfg=SimpleNamespace(volume_unit_label='стекове',driver_total_volume_label='Общо BGN',driver_customer_volume_label='лв.')
        self.assertEqual(volume_unit_label(cfg),'стекове')
        for short in (False,True):
            self.assertEqual(total_volume_unit_label(cfg,short=short),'Общо BGN')
            self.assertEqual(customer_volume_unit_label(cfg,short=short),'лв.')
        cfg.driver_total_volume_label=' '
        self.assertEqual(total_volume_unit_label(cfg,short=True),'ст.')
        self.assertEqual(customer_volume_unit_label(cfg),'лв.')
        cfg.driver_customer_volume_label=''
        cfg.volume_unit_label='палети'
        self.assertEqual(total_volume_unit_label(cfg),'палети')
        self.assertEqual(customer_volume_unit_label(cfg),'палети')

    def test_excel_uses_total_and_customer_units_without_changing_values(self):
        fixture=excel_tests.ExcelReportTests();fixture.setUp()
        try:
            fixture.exporter.config.volume_unit_label='fallback'
            fixture.exporter.config.driver_total_volume_label='палети'
            fixture.exporter.config.driver_customer_volume_label='кашони'
            wb=fixture.workbook()
            solution=SimpleNamespace(routes=[route([customer()])],total_distance_km=12)
            fixture.exporter._create_routes_sheet(wb,solution)
            fixture.exporter._create_summary_sheet(wb,solution,[])
            fixture.exporter._create_vehicle_stats_sheet(wb,solution)
            fixture.exporter._create_unserved_sheet(wb,[customer('u')])
            stream=io.BytesIO();wb.save(stream);stream.seek(0);saved=load_workbook(stream)
            self.assertEqual(fixture.by_header(saved['Маршрути'])['Обем (кашони)'].value,2.5)
            self.assertEqual(fixture.by_header(saved['Необслужени клиенти'])['Обем (кашони)'].value,2.5)
            stats=fixture.by_label(saved['Обобщение'])
            self.assertEqual(stats['Обслужен обем (палети)'],2.5)
            self.assertEqual(stats['Капацитет на активните бусове (палети)'],10)
            self.assertTrue(stats['Обща натовареност на бусовете (%)'].startswith('=B'))
            self.assertEqual(fixture.by_header(saved['Статистики по бусове'])['Общ обем (палети)'].value,2.5)
            self.assertEqual(_route_load_text(solution.routes[0],fixture.exporter.config),'2.5 палети')
            for sheet in saved:
                self.assertFalse(any('fallback' in str(cell.value) for row_cells in sheet for cell in row_cells))
        finally: fixture.doCleanups()

    def test_context_units_in_overview_maps_and_customer_popups(self):
        fixture=map_tests.VehicleMapGroupingTests
        config=fixture._config_without_zones()
        config.vehicles=[SimpleNamespace(vehicle_type='internal_bus',enabled=True,count=1,capacity=10)]
        current=fixture._route(1,'A',[fixture._customer('C','Client',42.71,'')])
        for provider in ('osm','google'):
            with self.subTest(provider=provider),patch('output_handler.get_config',return_value=config):
                generator=fixture._map_generator(provider)
                generator.config.volume_unit_label='fallback'
                generator.config.show_route_colors=True
                generator.config.driver_total_volume_label='Total <unit>'
                generator.config.driver_customer_volume_label='Client & unit'
                doc=generator.create_map(SimpleNamespace(routes=[current],dropped_customers=[]),SimpleNamespace(warehouse_customers=[]),(42.7,23.3))
                page=doc.html_content if provider=='google' else doc.get_root().render()
                if provider=='google':
                    data=json.loads(re.search(r'const MAP_DATA = (.*?);\s*let infoWindow;',page,re.S).group(1))
                    self.assertEqual(data['volumeUnit'],'Client & unit')
                    self.assertEqual(data['volumeUnitShort'],'Total <unit>')
                    self.assertEqual(data['routes'][0]['loadText'],'2.5 Total <unit>')
                    self.assertIn('Client &amp; unit',data['routes'][0]['markers'][0]['popup'])
                else:
                    decoded=page.replace(r'\u0026','&').replace(r'\u003c','<').replace(r'\u003e','>')
                    self.assertTrue('Total &lt;unit&gt;' in decoded,'Route filter should use the escaped total unit')
                    self.assertTrue('Client &amp; unit' in decoded,'Customer popup should use the escaped customer unit')
                    self.assertNotIn('Total <unit>',decoded)

    def test_csv_and_legacy_excel_exports_use_customer_unit(self):
        import csv
        fixture=excel_tests.ExcelReportTests();fixture.setUp()
        try:
            with tempfile.TemporaryDirectory() as directory:
                cfg=fixture.exporter.config
                cfg.excel_output_dir=directory;cfg.csv_output_file=str(Path(directory)/'routes.csv')
                cfg.driver_total_volume_label='палети';cfg.driver_customer_volume_label='кашони'
                solution=SimpleNamespace(routes=[route([customer()])])
                csv_path=fixture.exporter.export_routes_csv(solution)
                with open(csv_path,encoding='utf-8-sig',newline='') as file:
                    rows=list(csv.reader(file,delimiter=','))
                column=rows[0].index('Обем (кашони)')
                self.assertEqual(float(rows[1][column]),2.5)
                path=fixture.exporter.export_warehouse_orders([customer()])
                saved=load_workbook(path)
                self.assertEqual(fixture.by_header(saved.active)['Обем (кашони)'].value,2.5)
                path=fixture.exporter.export_vehicle_routes(solution)
                saved=load_workbook(path)
                self.assertEqual(fixture.by_header(saved.active)['Обем (кашони)'].value,2.5)
        finally: fixture.doCleanups()

    def test_charts_distinguish_route_totals_from_customer_quantities(self):
        import matplotlib
        matplotlib.use('Agg')
        from matplotlib.figure import Figure
        from output_handler import ChartGenerator
        cfg=OutputConfig(driver_total_volume_label='палети',driver_customer_volume_label='кашони')
        generator=ChartGenerator(cfg)
        solution=SimpleNamespace(routes=[route([customer()])])
        saved=[]
        with patch.object(Figure,'savefig',autospec=True,side_effect=lambda fig,*args,**kwargs:saved.append(fig)):
            generator._generate_route_comparison_chart(solution)
            generator._generate_volume_distribution_chart(solution)
        self.assertEqual(saved[0].axes[2].get_ylabel(),'палети')
        self.assertEqual(saved[1].axes[0].get_xlabel(),'Обем (кашони)')
        self.assertEqual(saved[1].axes[1].get_ylabel(),'Обем (кашони)')

    def test_desktop_capacity_and_customer_labels_refresh_separately(self):
        import config_gui
        gui=config_gui.ConfigGUI.__new__(config_gui.ConfigGUI)
        cfg=SimpleNamespace(volume_unit_label='fallback',driver_total_volume_label='палети',driver_customer_volume_label='кашони')
        gui.cfg=SimpleNamespace(output=cfg)
        gui.widgets={}
        seen={}
        def widget(key):
            return SimpleNamespace(winfo_exists=lambda:True,configure=lambda **values:seen.update({key:values['text']}))
        gui._quantity_labels=[(widget('capacity'),'Капацитет'),(widget('customer'),'Количество')]
        gui._refresh_volume_unit_labels()
        self.assertEqual(seen,{'capacity':'Капацитет (палети):','customer':'Количество (кашони):'})
        gui.widgets['output.driver_total_volume_label']=SimpleNamespace(get=lambda:'оборот')
        gui._refresh_volume_unit_labels()
        self.assertEqual(seen['capacity'],'Капацитет (оборот):')
        self.assertEqual(seen['customer'],'Количество (кашони):')

    def test_legacy_defaults_and_custom_labels(self):
        self.assertEqual(volume_unit_label(SimpleNamespace()), "стекове")
        self.assertEqual(volume_unit_label(SimpleNamespace(), short=True), "ст.")
        for label in ("кашони", "оборот", "лв.", "kg", "EUR"):
            cfg = OutputConfig(volume_unit_label=label)
            self.assertEqual(volume_unit_label(cfg), label)
            self.assertEqual(volume_unit_label(cfg, short=True), label)
        self.assertEqual(validate_volume_unit_label("  брой   кашони  "), "брой кашони")
        for invalid in ("", " " * 3, "X" * 41):
            with self.assertRaises(ValueError):
                validate_volume_unit_label(invalid)

    def test_web_legacy_default_override_and_global_field_insertion(self):
        base = deepcopy(get_config())
        base.output = SimpleNamespace(**{k:v for k,v in vars(base.output).items() if k != "volume_unit_label"})
        self.assertEqual(api._web_run_settings_from_config(base)["output"]["volume_unit_label"], "стекове")
        with patch.object(api, "get_config", return_value=base):
            cfg, applied, ignored = api._build_web_run_config_override({"run_settings":{"output":{"volume_unit_label":"кашони"}}})
        self.assertEqual(cfg.output.volume_unit_label, "кашони")
        self.assertFalse(hasattr(base.output, "volume_unit_label"))
        self.assertIn("output.volume_unit_label", applied)
        self.assertEqual(ignored, [])
        source = 'from dataclasses import dataclass\n@dataclass\nclass OutputConfig:\n    enable_interactive_map: bool = True\n'
        for label in ('кашони', 'оборот "лв."'):
            source = api._replace_class_field_literal(source, "OutputConfig", "volume_unit_label", label)
            scope = {}
            exec(compile(source, "test_config.py", "exec"), scope)
            self.assertEqual(scope["OutputConfig"]().volume_unit_label, label)
        self.assertEqual(source.count("volume_unit_label:"), 1)

    def test_excel_custom_units_remain_numeric_and_formula_linked(self):
        fixture = excel_tests.ExcelReportTests()
        fixture.setUp()
        try:
            fixture.exporter.config.volume_unit_label = "кашони"
            wb = fixture.workbook()
            solution = SimpleNamespace(routes=[route([customer()])], total_distance_km=12)
            fixture.exporter._create_routes_sheet(wb, solution)
            fixture.exporter._create_summary_sheet(wb, solution, [])
            fixture.exporter._create_vehicle_stats_sheet(wb, solution)
            fixture.exporter._create_unserved_sheet(wb, [customer("u")])
            stream = io.BytesIO()
            wb.save(stream)
            stream.seek(0)
            saved = load_workbook(stream)
            self.assertEqual(fixture.by_header(saved["Маршрути"])["Обем (кашони)"].value, 2.5)
            stats = fixture.by_label(saved["Обобщение"])
            self.assertEqual(stats["Обслужен обем (кашони)"], 2.5)
            self.assertEqual(stats["Капацитет на активните бусове (кашони)"], 10)
            self.assertTrue(stats["Обща натовареност на бусовете (%)"].startswith("=B"))
            self.assertEqual(_route_load_text(solution.routes[0], fixture.exporter.config), "2.5 кашони")
            self.assertEqual(fixture.by_header(saved["Статистики по бусове"])["Капацитет използване (%)"].value, .25)
            for ws in saved:
                self.assertFalse(any("(ст.)" in str(cell.value) for cells in ws for cell in cells))
        finally:
            fixture.doCleanups()

    def test_empty_unserved_controls_and_safe_custom_units_in_both_maps(self):
        fixture = map_tests.VehicleMapGroupingTests
        config = fixture._config_without_zones()
        config.vehicles = [SimpleNamespace(vehicle_type="internal_bus", enabled=True, count=1, capacity=10)]
        client = fixture._customer("C", "Client", 42.71, "")
        current = fixture._route(1, "A", [client])
        solution = SimpleNamespace(routes=[current], dropped_customers=[])
        allocation = SimpleNamespace(warehouse_customers=[])
        for provider in ("osm", "google"):
            with self.subTest(provider=provider), patch("output_handler.get_config", return_value=config):
                generator = fixture._map_generator(provider)
                generator.config.volume_unit_label = '<unit "&">'
                document = generator.create_map(solution, allocation, (42.7, 23.3))
                page = document.html_content if provider == "google" else document.get_root().render()
                self.assertNotIn("25.0%", page)
                self.assertIn("Клиенти на картата", page)
                self.assertIn("route-load", page)
                if provider == "osm":
                    self.assertNotIn("UnservedCustomerControl", [type(child).__name__ for child in document._children.values()])
                    self.assertIn("&lt;unit &quot;&amp;&quot;&gt;", page)
                    self.assertNotIn('<unit "&">', page)
                else:
                    data = json.loads(re.search(r"const MAP_DATA = (.*?);\s*let infoWindow;", page, re.S).group(1))
                    self.assertEqual(data["unserved"]["total"], 0)
                    self.assertIn("MAP_DATA.unserved.total > 0", page)
                    self.assertEqual(data["volumeUnit"], '<unit "&">')
                    self.assertEqual(data["routes"][0]["loadText"], '2.5 <unit "&">')

    def test_desktop_save_adds_unit_field_to_legacy_config_only_in_temp_directory(self):
        import config_gui
        gui = config_gui.ConfigGUI.__new__(config_gui.ConfigGUI)
        gui.cfg = deepcopy(get_config())
        gui.cfg.vehicles = []
        source = 'from dataclasses import dataclass\n@dataclass\nclass OutputConfig:\n    enable_interactive_map: bool = True\n'
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "config.py"
            path.write_text(source, encoding="utf-8")
            with patch.object(config_gui, "_base_dir", directory):
                gui._apply_to_config_file({"output.volume_unit_label": "оборот"})
            scope = {}
            exec(compile(path.read_text(encoding="utf-8"), "saved_config.py", "exec"), scope)
            self.assertEqual(scope["OutputConfig"]().volume_unit_label, "оборот")
            self.assertTrue(scope["OutputConfig"]().enable_interactive_map)


if __name__ == "__main__":
    unittest.main()
