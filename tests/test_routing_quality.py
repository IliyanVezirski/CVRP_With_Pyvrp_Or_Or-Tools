"""Synthetic routing provenance tests. No real routing, uploads or optimisation."""
import io
import json
import pickle
import tempfile
import unittest
from contextlib import redirect_stderr, redirect_stdout
from copy import deepcopy
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock, patch

import requests
from openpyxl import load_workbook

from config import OSRMConfig, OutputConfig, get_config
from osrm_client import DistanceMatrix, OSRMCache, OSRMClient
from routing_quality import (
    OSRM, UNKNOWN, VALHALLA, RoutingResponseError, matrix_quality_counts,
    quality_at, quality_code, quality_grid, quality_kind, route_quality_edges,
)
from valhalla_client import ValhallaClient
from tests import test_excel_report as excel_tests

customer, route = excel_tests.customer, excel_tests.route


LOCATIONS = [(42.0, 23.0), (42.01, 23.01), (43.0, 23.5)]


def matrix(locations=LOCATIONS, code=OSRM):
    n = len(locations)
    return DistanceMatrix(
        [[0 if i == j else 1234.5 for j in range(n)] for i in range(n)],
        [[0 if i == j else 120 for j in range(n)] for i in range(n)],
        list(locations), list(range(n)), list(range(n)), quality_grid(n, value=code),
    )


class RoutingQualityTests(unittest.TestCase):
    def setUp(self):
        self.client = OSRMClient.__new__(OSRMClient)
        self.client.config = OSRMConfig(base_url="https://example.invalid", average_speed_kmh=40)
        self.client.cache = None
        self.client.session = Mock()
        self.no_progress = redirect_stderr(io.StringIO())
        self.no_progress.__enter__()
        self.addCleanup(self.no_progress.__exit__, None, None, None)

    def test_table_normalisation_distinguishes_sources_and_nulls(self):
        distances, durations, quality = self.client._normalise_table_values({
            "distances": [[0, None, 400], [1000, 0, 300], [None, 500, 0]],
            "durations": [[0, None, None], [100, 0, 30], [60, 50, 0]],
        }, LOCATIONS)
        self.assertEqual(quality[0][1], quality_code("approx", "approx", "no_route"))
        self.assertGreater(distances[0][1], 0)
        self.assertGreater(durations[0][1], 0)
        self.assertEqual(quality[0][2], quality_code("osrm", "approx", "invalid_value"))
        self.assertEqual(distances[0][2], 400)
        self.assertEqual(quality[2][0], quality_code("derived", "osrm", "distance_from_duration"))
        self.assertEqual(durations[2][0], 60)
        self.assertEqual(distances[2][0], 60 * 40 * 1000 / 3600 * 1.1)
        self.assertEqual(quality[1][0], OSRM)
        self.assertEqual(distances[1][0], 1000)

    def test_invalid_values_and_server_estimates_are_not_marked_as_road(self):
        _, _, quality = self.client._normalise_table_values({
            "distances": [[0, float("nan")], [120, 0]],
            "durations": [[0, -3], [12, 0]], "fallback_speed_cells": [[1, 0]],
        }, LOCATIONS[:2])
        self.assertEqual(quality[0][1], quality_code("approx", "approx", "invalid_value"))
        self.assertEqual(quality[1][0], quality_code("approx", "approx", "server_estimate"))
        _, _, quality = self.client._normalise_table_values({"distances": [], "durations": []}, LOCATIONS[:2])
        self.assertEqual(quality[0][1], quality_code("approx", "approx", "invalid_value"))

    def test_table_and_route_response_paths_attach_quality(self):
        response = Mock()
        response.json.return_value = {"code": "Ok", "distances": [[0, 10], [20, 0]], "durations": [[0, 1], [2, 0]]}
        self.client.session.get.return_value = response
        with patch.object(self.client, "_approaches_query", return_value=""):
            actual = self.client._try_table_api_to_server(LOCATIONS[:2], self.client.config.base_url)
        self.assertEqual(actual.quality[0][1], OSRM)
        with patch.object(self.client, "_single_route_request", side_effect=[
            requests.exceptions.Timeout("secret URL must not be stored"), {"distance": 20, "duration": 2}
        ]):
            actual = self.client._build_matrix_via_routes_only(LOCATIONS[:2])
        self.assertEqual(actual.quality[0][1], quality_code("approx", "approx", "timeout"))
        self.assertEqual(actual.quality[1][0], OSRM)
        self.assertNotIn("secret", str(actual.quality))

    def test_no_route_is_distinct_from_a_failed_request(self):
        self.client.session.get.return_value.json.return_value = {"code": "NoRoute"}
        with patch.object(self.client, "_build_route_url", return_value="https://example.invalid"):
            with self.assertRaises(RoutingResponseError) as caught:
                self.client._single_route_request(*LOCATIONS[:2])
        self.assertEqual(caught.exception.reason, "no_route")

    def test_optimized_batches_copy_quality_in_each_direction(self):
        def request(locations):
            result = matrix(locations)
            result.quality[0][-1] = quality_code("approx", "approx", "timeout")
            return result
        with patch.object(self.client, "_try_table_api_with_fallback", side_effect=request):
            self.client.config.max_locations_for_osrm = 4
            actual = self.client._build_optimized_table_batches(LOCATIONS)
        self.assertEqual(actual.quality[0][1], quality_code("approx", "approx", "timeout"))
        self.assertEqual(actual.quality[0][2], quality_code("approx", "approx", "timeout"))
        self.assertEqual(actual.quality[2][0], OSRM)

    def test_failed_batch_marks_every_replaced_cell(self):
        with patch.object(self.client, "_try_table_api_with_fallback", side_effect=requests.exceptions.Timeout()):
            actual = self.client._build_optimized_table_batches(LOCATIONS[:2])
        self.assertEqual(actual.quality[0][1], quality_code("approx", "approx", "timeout"))

    def test_interbatch_failure_preserves_successful_reverse_direction(self):
        actual = matrix(LOCATIONS[:2])
        with (patch.object(self.client, "_try_table_api_with_fallback", side_effect=requests.exceptions.Timeout()),
              patch.object(self.client, "_single_route_request", side_effect=[
                  RoutingResponseError("no route", "no_route"), {"distance": 765, "duration": 87}
              ])):
            self.client._fill_inter_batch_connections(actual.locations, actual.distances, actual.durations,
                                                      [(0, 1), (1, 2)], actual.quality)
        self.assertEqual(actual.quality[0][1], quality_code("approx", "approx", "no_route"))
        self.assertEqual(actual.quality[1][0], OSRM)
        self.assertEqual(actual.distances[1][0], 765)

    def test_legacy_builders_do_not_claim_estimates_are_road_values(self):
        with patch.object(self.client, "_try_table_api_with_fallback", return_value=matrix(LOCATIONS[:2])):
            actual = self.client._build_matrix_via_small_batches(LOCATIONS[:2])
        self.assertEqual(actual.quality[0][1], OSRM)
        with patch.object(self.client, "_single_route_request", side_effect=requests.exceptions.Timeout()):
            actual = matrix(LOCATIONS[:2])
            self.client._fill_batch_with_routes(LOCATIONS[:2], LOCATIONS[:2], actual.distances,
                                                actual.durations, 0, 0, actual.quality)
        self.assertEqual(actual.quality[0][1], quality_code("approx", "approx", "timeout"))
        approximate = self.client._create_approximate_matrix(LOCATIONS, "size_limit")
        self.assertEqual(matrix_quality_counts(approximate)["approximate"], 6)
        matching = self.client._process_match_api_response({"matchings": [{"legs": [{"distance": 5, "duration": 2}]}]}, LOCATIONS[:2])
        self.assertEqual(quality_kind(matching.quality[1][0]), "approximate")

    def test_cache_pickle_and_submatrix_preserve_provenance_and_legacy_unknown(self):
        actual = matrix()
        actual.quality[2][1] = quality_code("approx", "approx", "timeout")
        with tempfile.TemporaryDirectory() as directory:
            cache = OSRMCache(str(Path(directory) / "matrix.json"))
            cache.set(actual)
            reloaded = OSRMCache(cache.cache_file)
            self.assertEqual(reloaded.get(actual.locations).quality, actual.quality)
            central = reloaded.get_complete_central_matrix()
            submatrix = reloaded.extract_submatrix(central, [LOCATIONS[2], LOCATIONS[1]])
            self.assertEqual(submatrix.quality[0][1], actual.quality[2][1])
            self.assertEqual(pickle.loads(pickle.dumps(submatrix)).quality, submatrix.quality)
            for entry in reloaded.cache_data.values():
                entry.pop("quality")
            old = reloaded.get(actual.locations)
            self.assertEqual(quality_at(old, 2, 1), UNKNOWN)
            self.assertEqual(matrix_quality_counts(old)["unknown"], 6)

    def test_valhalla_partial_cells_track_distance_and_time_separately(self):
        client = ValhallaClient.__new__(ValhallaClient)
        client._recover_matrix_cell = lambda *args: (None, "invalid_value")
        distances, durations, quality, count = client._matrix_values_with_quality({
            "sources_to_targets": [[{"distance": 1.25, "time": None}], [{"distance": None, "time": 60}]]
        }, LOCATIONS[:2], LOCATIONS[2:], {"costing": "auto"})
        self.assertEqual(distances[0][0], 1250)
        self.assertEqual(durations[1][0], 60)
        self.assertEqual(quality[0][0], quality_code("valhalla", "approx", "invalid_value"))
        self.assertEqual(quality[1][0], quality_code("approx", "valhalla", "invalid_value"))
        self.assertEqual(count, 2)


class RoutingQualityReportTests(unittest.TestCase):
    def setUp(self):
        self.fixture = excel_tests.ExcelReportTests()
        self.fixture.setUp()
        self.addCleanup(self.fixture.doCleanups)
        self.client = customer(name="=Example", coordinates=LOCATIONS[1])
        self.actual = matrix([LOCATIONS[0], LOCATIONS[2], LOCATIONS[1]])
        self.solution = SimpleNamespace(routes=[route([self.client], depot_location=LOCATIONS[2],
                                                       end_location=LOCATIONS[0])], total_distance_km=20)

    def test_disabled_option_leaves_report_unchanged(self):
        wb = self.fixture.workbook()
        self.fixture.exporter._create_summary_sheet(wb, self.solution, [])
        with tempfile.TemporaryDirectory() as directory:
            self.fixture.exporter.config.excel_output_dir = directory
            path = self.fixture.exporter.export_all_to_single_excel(self.solution, [], self.actual, [self.client])
            saved = load_workbook(path)
            try:
                self.assertNotIn("Приблизителни отсечки", saved.sheetnames)
                self.assertNotIn("Отсечки с приблизителни данни", self.fixture.by_label(saved["Обобщение"]))
            finally:
                saved.close()

    def test_report_lists_only_selected_problem_edges_with_safe_text_and_ids(self):
        self.actual.quality[1][2] = quality_code("approx", "approx", "timeout")
        self.actual.quality[0][2] = quality_code("approx", "approx", "no_route")  # Not selected.
        self.fixture.exporter.config.enable_routing_quality_report = True
        with tempfile.TemporaryDirectory() as directory:
            self.fixture.exporter.config.excel_output_dir = directory
            path = self.fixture.exporter.export_all_to_single_excel(self.solution, [], self.actual, [self.client])
            saved = load_workbook(path)
            try:
                ws = saved["Приблизителни отсечки"]
                cells = self.fixture.by_header(ws)
                self.assertEqual(ws.max_row, 2)
                self.assertEqual(cells["ID до клиент"].value, "00001")
                self.assertEqual(cells["До"].value, "=Example")
                self.assertEqual(cells["До"].data_type, "s")
                self.assertEqual(cells["GPS от"].value, "43.000000, 23.500000")
                self.assertEqual(cells["Разстояние по матрица (км)"].value, 1.2345)
                self.assertEqual(cells["Пътуване по матрица (мин)"].value, 2)
                self.assertEqual(cells["Причина"].value, "Сървърът не отговори навреме")
                self.assertEqual(ws.freeze_panes, "A2")
                self.assertEqual(ws.auto_filter.ref, ws.dimensions)
                stats = self.fixture.by_label(saved["Обобщение"])
                self.assertEqual(stats["Отсечки с приблизителни данни"], 1)
                self.assertEqual(stats["Приблизителни връзки в цялата матрица"], 2)
                self.assertTrue(stats["Обща натовареност на бусовете (%)"].startswith("=B"))
            finally:
                saved.close()

    def test_clean_unknown_and_multiple_trip_reports_include_return_to_depot(self):
        self.solution.routes.append(route([self.client], trip_number=2))
        self.actual.quality[2][0] = UNKNOWN
        edges = list(route_quality_edges(self.solution, self.actual, [self.client]))
        self.assertEqual(len(edges), 4)
        self.assertEqual([e["route"].trip_number for e in edges], [1, 1, 2, 2])
        self.assertEqual(edges[1]["to_name"], "Депо / крайна точка")
        self.assertEqual(edges[1]["quality"], UNKNOWN)
        wb = self.fixture.workbook()
        self.fixture.exporter._create_summary_sheet(wb, self.solution, [])
        self.actual.quality = quality_grid(3, value=OSRM)
        self.fixture.exporter._add_routing_quality_report(wb, self.solution, self.actual, [self.client])
        self.assertNotIn("Приблизителни отсечки", wb.sheetnames)
        self.assertEqual(self.fixture.by_label(wb["Обобщение"])["Отсечки с пътни данни"], 4)

    def test_missing_matrix_and_ambiguous_visits_never_claim_road_data(self):
        edges = list(route_quality_edges(self.solution, None, [self.client]))
        self.assertTrue(all(quality_kind(e["quality"]) == "unknown" for e in edges))
        duplicate = matrix([LOCATIONS[0], LOCATIONS[2], LOCATIONS[1], LOCATIONS[1]])
        edges = list(route_quality_edges(self.solution, duplicate, [self.client, deepcopy(self.client)]))
        self.assertTrue(all(quality_kind(e["quality"]) == "unknown" for e in edges))

    def test_output_handler_passes_the_original_matrix_and_order_without_network(self):
        from output_handler import OutputHandler
        cfg = self.fixture.exporter.config
        cfg.enable_interactive_map = False
        cfg.enable_csv_output = False
        cfg.enable_charts = False
        cfg.enable_routing_quality_report = True
        allocation = SimpleNamespace(vehicle_customers=[self.client], warehouse_customers=[])
        self.solution.dropped_customers = []
        handler = OutputHandler(cfg)
        with (patch.object(handler.excel_exporter, "export_all_to_single_excel", return_value="synthetic.xlsx") as export,
              patch("requests.sessions.Session.request", side_effect=AssertionError("No network"))):
            result = handler.generate_all_outputs(self.solution, allocation, LOCATIONS[0], self.actual)
        self.assertEqual(result["excel_report"], "synthetic.xlsx")
        self.assertIs(export.call_args.kwargs["distance_matrix"], self.actual)
        self.assertIs(export.call_args.kwargs["matrix_customers"], allocation.vehicle_customers)

    def test_main_output_boundary_retains_matrix_provenance(self):
        import main
        with (patch.object(main, "_assert_mandatory_customer_solution"),
              patch.object(main, "OutputHandler") as handler,
              patch.object(main, "_print_summary")):
            main.process_results(self.solution, SimpleNamespace(depot_location=LOCATIONS[0]),
                                 SimpleNamespace(), 1, LOCATIONS[:2], distance_matrix=self.actual)
        self.assertIs(handler.return_value.generate_all_outputs.call_args.kwargs["distance_matrix"], self.actual)


class RoutingQualitySettingsTests(unittest.TestCase):
    def test_legacy_default_web_override_and_field_insertion(self):
        import cvrp_api_server as api
        base = deepcopy(get_config())
        base.output = SimpleNamespace(**{key: value for key, value in vars(base.output).items()
                                         if key != "enable_routing_quality_report"})
        self.assertFalse(api._web_run_settings_from_config(base)["output"]["enable_routing_quality_report"])
        with patch.object(api, "get_config", return_value=base):
            config, applied, ignored = api._build_web_run_config_override({
                "run_settings": {"output": {"enable_routing_quality_report": True}}})
        self.assertTrue(config.output.enable_routing_quality_report)
        self.assertIn("output.enable_routing_quality_report", applied)
        self.assertEqual(ignored, [])
        self.assertFalse(hasattr(base.output, "enable_routing_quality_report"))
        source = 'from dataclasses import dataclass\n@dataclass\nclass OutputConfig:\n    enable_excel_output: bool = True\n'
        for value in (True, False):
            source = api._replace_class_field_literal(source, "OutputConfig", "enable_routing_quality_report", value)
            scope = {}
            exec(compile(source, "synthetic_config.py", "exec"), scope)
            self.assertEqual(scope["OutputConfig"]().enable_routing_quality_report, value)
        self.assertEqual(source.count("enable_routing_quality_report:"), 1)

    def test_desktop_save_inserts_option_into_an_old_config(self):
        import config_gui
        gui = config_gui.ConfigGUI.__new__(config_gui.ConfigGUI)
        gui.cfg = SimpleNamespace(vehicles=[])
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "config.py"
            path.write_text('from dataclasses import dataclass\n@dataclass\nclass OutputConfig:\n    enable_excel_output: bool = True\n', encoding="utf-8")
            with patch.object(config_gui, "_base_dir", directory):
                self.assertTrue(gui._apply_to_config_file({"output.enable_routing_quality_report": True}))
                self.assertIn("enable_routing_quality_report: bool = True", path.read_text(encoding="utf-8"))
                self.assertTrue(gui._apply_to_config_file({"output.enable_routing_quality_report": False}))
                self.assertIn("enable_routing_quality_report: bool = False", path.read_text(encoding="utf-8"))


if __name__ == "__main__":
    unittest.main()
