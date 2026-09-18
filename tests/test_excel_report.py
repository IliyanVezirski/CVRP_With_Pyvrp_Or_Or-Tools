"""Report regressions using synthetic data only: no optimisation or HTTP calls."""
import csv
import io
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from config import OutputConfig, VehicleType
from input_handler import Customer
from output_handler import ExcelExporter, InteractiveMapGenerator


def customer(identifier="00001", **kwargs):
    values = dict(id=identifier, name="Client", coordinates=(42.01, 23.01),
                  volume=2.5, original_gps_data="42.01,23.01", document="0000123456")
    values.update(kwargs)
    return Customer(**values)


def route(customers=None, **kwargs):
    clients = list(customers or [])
    values = dict(vehicle_type=VehicleType.INTERNAL_BUS, vehicle_id=0,
                  vehicle_name="Van 1", vehicle_key="van-1", trip_number=1,
                  customers=clients, depot_location=(42.0, 23.0), end_location=(42.0, 23.0),
                  total_volume=sum(c.volume for c in clients), total_distance_km=12.34567,
                  total_time_minutes=30.125, planned_start_minutes=480)
    values.update(kwargs)
    return SimpleNamespace(**values)


class ExcelReportTests(unittest.TestCase):
    def setUp(self):
        self.vehicle = SimpleNamespace(vehicle_type=VehicleType.INTERNAL_BUS, enabled=True,
                                       count=1, name="Van 1", capacity=10,
                                       service_time_minutes=15, start_time_minutes=480)
        self.config = SimpleNamespace(
            vehicles=[self.vehicle],
            cvrp=SimpleNamespace(enable_customer_time_windows=False, global_start_time_minutes=480),
            locations=SimpleNamespace(center_location=(42.0, 23.0)),
        )
        self.config_patch = patch("output_handler.get_config", return_value=self.config)
        self.config_patch.start()
        self.addCleanup(self.config_patch.stop)
        self.exporter = ExcelExporter(OutputConfig())

    @staticmethod
    def workbook():
        wb = Workbook()
        wb.remove(wb.active)
        return wb

    @staticmethod
    def by_label(ws):
        return {row[0].value: row[1].value for row in ws.iter_rows()}

    @staticmethod
    def by_header(ws, row=2):
        return {cell.value: ws.cell(row=row, column=cell.column) for cell in ws[1]}

    def test_summary_counts_visits_not_routes_and_labels_served_volume(self):
        routes = [route([customer("a"), customer("b")]), route([customer("c")], vehicle_key="van-2")]
        solution = SimpleNamespace(routes=routes, total_distance_km=24.69134, total_time_minutes=60.25)
        wb = self.workbook()
        self.exporter._create_summary_sheet(wb, solution, [customer("d")])
        stats = self.by_label(wb["Обобщение"])
        self.assertEqual(stats["Общо клиенти"], 4)
        self.assertEqual(stats["Обслужени клиенти"], 3)
        self.assertEqual(stats["Необслужени клиенти"], 1)
        self.assertEqual(stats["Обслужен обем (ст.)"], 7.5)
        self.assertEqual(stats["Общо разстояние (км)"], 24.69134)
        self.assertAlmostEqual(stats["Общо време (ч)"], 60.25 / 60)

    def test_summary_empty_plan(self):
        wb = self.workbook()
        self.exporter._create_summary_sheet(wb, SimpleNamespace(routes=[], total_distance_km=0), [])
        stats = self.by_label(wb["Обобщение"])
        self.assertEqual(stats["Общо клиенти"], 0)
        self.assertEqual(stats["Общо време (ч)"], 0)
        self.assertEqual(stats["Капацитет на активните бусове (ст.)"], 0)
        self.assertIsNone(stats["Обща натовареност на бусовете (%)"])
        self.assertEqual(stats["Брой физически бусове"], 0)
        self.assertEqual(stats["Обслужени заявки"], 0)
        self.assertIsNone(stats["Среден брой заявки на бус"])

    def test_summary_utilisation_counts_each_active_physical_vehicle_once(self):
        large_vehicle = SimpleNamespace(**vars(self.vehicle))
        large_vehicle.name = "Van 2"
        large_vehicle.capacity = 30
        self.config.vehicles.append(large_vehicle)
        # Inactive vehicles and empty routes do not dilute the active-fleet ratio.
        self.config.vehicles.append(SimpleNamespace(**{**vars(self.vehicle), "capacity": 1000}))
        routes = [
            route([customer("a", volume=8)]),
            route([customer("b", volume=2)], trip_number=2),
            route([customer("c", volume=15)], vehicle_id=1, vehicle_key="van-2", vehicle_name="Van 2"),
            route([], vehicle_id=2, vehicle_key="unused"),
        ]
        wb = self.workbook()
        self.exporter._create_summary_sheet(wb, SimpleNamespace(routes=routes, total_distance_km=0), [customer("u", volume=100)])
        ws = wb["Обобщение"]
        cells = {row[0].value: row[1] for row in ws.iter_rows()}
        capacity_cell = cells["Капацитет на активните бусове (ст.)"]
        volume_cell = cells["Обслужен обем (ст.)"]
        usage_cell = cells["Обща натовареност на бусовете (%)"]
        self.assertEqual(capacity_cell.value, 40)
        self.assertEqual(volume_cell.value, 25)
        self.assertEqual(usage_cell.value, f"={volume_cell.coordinate}/{capacity_cell.coordinate}")
        self.assertEqual(usage_cell.number_format, "0.0%")
        self.assertEqual(volume_cell.value / capacity_cell.value, 0.625)
        self.assertEqual(cells["Брой физически бусове"].value, 2)

    def test_reload_can_produce_150_percent_without_overloading_a_course(self):
        self.vehicle.capacity = 400
        routes = [route([customer("a", volume=400)]),
                  route([customer("b", volume=200)], trip_number=2)]
        wb = self.workbook()
        self.exporter._create_summary_sheet(wb, SimpleNamespace(routes=routes, total_distance_km=0), [])
        stats = self.by_label(wb["Обобщение"])
        self.assertEqual(stats["Брой физически бусове"], 1)
        self.assertEqual(stats["Обслужен обем (ст.)"] / stats["Капацитет на активните бусове (ст.)"], 1.5)

    def test_orders_average_uses_documents_and_physical_vehicles_not_courses(self):
        routes = [route([customer("a", grouped_documents=[{"document": "001"}, {"document": "002"}])]),
                  route([customer("b")], trip_number=2),
                  route([customer("c")], vehicle_key="van-2"),
                  route([], vehicle_key="unused")]
        wb = self.workbook()
        self.exporter._create_summary_sheet(wb, SimpleNamespace(routes=routes, total_distance_km=0), [customer("unserved")])
        cells = {row[0].value: row[1] for row in wb["Обобщение"].iter_rows()}
        orders, vehicles, average = [cells[label] for label in ("Обслужени заявки", "Брой физически бусове", "Среден брой заявки на бус")]
        self.assertEqual(orders.value, 4)
        self.assertEqual(vehicles.value, 2)
        self.assertEqual(average.value, f"={orders.coordinate}/{vehicles.coordinate}")
        self.assertEqual(average.number_format, "0.0")
        self.assertEqual(orders.value / vehicles.value, 2)

    def test_same_display_names_and_legacy_routes_are_not_merged(self):
        for keys in (("van-1", "van-2"), (None, None)):
            with self.subTest(keys=keys):
                routes = [route([customer("a")], vehicle_key=keys[0]),
                          route([customer("b")], vehicle_key=keys[1])]
                self.assertEqual(self.exporter._total_vehicle_capacity_for_report(routes), 20)

    def test_summary_unknown_capacity_does_not_produce_a_partial_percentage(self):
        routes = [route([customer()]), route([customer("b")], vehicle_id=1)]
        for invalid_capacity in (None, 0, -1, float("nan"), float("inf"), "invalid"):
            with self.subTest(capacity=invalid_capacity):
                self.config.vehicles = [self.vehicle, SimpleNamespace(**{**vars(self.vehicle), "capacity": invalid_capacity})]
                wb = self.workbook()
                self.exporter._create_summary_sheet(wb, SimpleNamespace(routes=routes, total_distance_km=0), [])
                stats = self.by_label(wb["Обобщение"])
                self.assertIsNone(stats["Капацитет на активните бусове (ст.)"])
                self.assertIsNone(stats["Обща натовареност на бусовете (%)"])

    def test_summary_includes_reload_and_other_vehicle_time(self):
        routes = [route([customer("a")], planned_start_minutes=480, planned_end_minutes=540, total_time_minutes=60),
                  route([customer("b")], trip_number=2, planned_start_minutes=570, planned_end_minutes=630, total_time_minutes=60),
                  route([customer("c")], vehicle_key="van-2", total_time_minutes=30)]
        wb = self.workbook()
        self.exporter._create_summary_sheet(wb, SimpleNamespace(routes=routes, total_distance_km=0), [])
        stats = self.by_label(wb["Обобщение"])
        self.assertEqual(stats["Общо време (ч)"], 3)
        self.assertEqual(stats["Брой физически бусове"], 2)
        self.assertEqual(stats["Брой курсове"], 3)

    def test_customer_text_is_literal_and_leading_zero_ids_survive_save(self):
        client = customer(name='=HYPERLINK("https://example.invalid", "test")',
                          delivery_comment="=1+1")
        wb = self.workbook()
        self.exporter._create_routes_sheet(wb, SimpleNamespace(routes=[route([client])]))
        self.exporter._create_unserved_sheet(wb, [client])
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        reloaded = load_workbook(buffer)
        values = self.by_header(reloaded["Маршрути"])
        self.assertEqual(values["Име клиент"].value, client.name)
        self.assertEqual(values["Име клиент"].data_type, "s")
        self.assertEqual(values["Коментар доставка"].data_type, "s")
        self.assertEqual(values["Номер поръчка"].value, "0000123456")
        self.assertEqual(values["ID клиент"].value, "00001")
        self.assertEqual(values["Номер поръчка"].number_format, "@")
        self.assertEqual(reloaded["Необслужени клиенти"]["B2"].data_type, "s")

    def test_report_precision_capacity_and_table_navigation(self):
        wb = self.workbook()
        current_route = route([customer()])
        solution = SimpleNamespace(routes=[current_route])
        self.exporter._create_routes_sheet(wb, solution)
        self.exporter._create_unserved_sheet(wb, [customer("u")])
        self.exporter._create_vehicle_stats_sheet(wb, solution)
        cells = self.by_header(wb["Статистики по бусове"])
        self.assertEqual(cells["Капацитет използване (%)"].value, 0.25)
        self.assertEqual(cells["Капацитет използване (%)"].number_format, "0.0%")
        self.assertEqual(cells["Разстояние (км)"].value, 12.34567)
        self.assertEqual(cells["Разстояние (км)"].number_format, "0.00")
        self.assertAlmostEqual(cells["Време (ч)"].value, 30.125 / 60)
        for ws in wb.worksheets:
            self.assertEqual(ws.freeze_panes, "A2")
            self.assertEqual(ws.auto_filter.ref, ws.dimensions)
            for header in ws[1]:
                self.assertFalse(header.alignment.wrap_text)
                self.assertGreaterEqual(ws.column_dimensions[header.column_letter].width, len(str(header.value or "")) * 1.15 + 5)

    def test_unknown_capacity_is_not_a_false_zero_without_center_columns(self):
        missing = customer("missing", coordinates=None)
        valid = customer("valid")
        self.vehicle.capacity = 0
        wb = self.workbook()
        self.exporter._create_unserved_sheet(wb, [missing])
        self.assertEqual(list(self.by_header(wb["Необслужени клиенти"])),
                         ["ID", "Име", "Обем (ст.)", "GPS координати"])
        self.exporter._create_vehicle_stats_sheet(wb, SimpleNamespace(routes=[route([missing, valid])]))
        cells = self.by_header(wb["Статистики по бусове"])
        self.assertIsNone(cells["Капацитет използване (%)"].value)
        self.assertFalse(any("центъра" in header.lower() for header in cells))
        self.assertEqual(cells["Депо стартова точка"].value, "42.000000, 23.000000")

    def test_long_custom_quantity_heading_gets_full_single_line_width(self):
        unit = "максимално дълга мерна единица за тест"
        self.exporter.config.volume_unit_label = unit
        wb = self.workbook()
        self.exporter._create_vehicle_stats_sheet(wb, SimpleNamespace(routes=[route([customer()])]))
        header = next(cell for cell in wb["Статистики по бусове"][1] if cell.value == f"Общ обем ({unit})")
        self.assertFalse(header.alignment.wrap_text)
        self.assertGreater(wb["Статистики по бусове"].column_dimensions[header.column_letter].width, 50)

    def test_compact_rows_and_removed_columns_survive_save_with_multiple_courses(self):
        long_name = "Дълго име на клиент " * 6
        clients = [customer("a"), customer("b", name=long_name, delivery_comment="Първи ред\nВтори ред")]
        for multiple_courses in (False, True):
            with self.subTest(multiple_courses=multiple_courses):
                routes = [route(clients)]
                if multiple_courses:
                    routes.append(route([customer("c")], trip_number=2))
                wb = self.workbook()
                self.exporter._create_routes_sheet(wb, SimpleNamespace(routes=routes))
                self.exporter._create_unserved_sheet(wb, clients)
                self.exporter._create_vehicle_stats_sheet(wb, SimpleNamespace(routes=routes))
                buffer = io.BytesIO()
                wb.save(buffer)
                buffer.seek(0)
                saved = load_workbook(buffer)
                for ws in saved.worksheets:
                    self.assertFalse(any("центъра" in str(cell.value).lower() for cell in ws[1]))
                    self.assertEqual(ws.row_dimensions[1].height, 26)
                    for cells in ws.iter_rows(min_row=2):
                        self.assertEqual(ws.row_dimensions[cells[0].row].height, 15)
                        self.assertTrue(all(not cell.alignment.wrap_text for cell in cells))
                    self.assertEqual(ws.auto_filter.ref, ws.dimensions)
                headers = self.by_header(saved["Маршрути"], row=3)
                self.assertNotIn("До центъра по права линия (км)", headers)
                self.assertEqual(headers["Име клиент"].value, long_name)
                self.assertEqual(headers["Коментар доставка"].value, "Първи ред\nВтори ред")
                self.assertEqual(headers["Депо стартова точка"].value, "42.000000, 23.000000")
                self.assertEqual("Курс" in headers, multiple_courses)
                unserved = self.by_header(saved["Необслужени клиенти"])
                self.assertNotIn("Latitude", unserved)
                self.assertNotIn("Longitude", unserved)
                self.assertEqual(unserved["GPS координати"].value, clients[0].original_gps_data)
                self.assertEqual(saved["Необслужени клиенти"].max_column, 4)
                stats = self.by_header(saved["Статистики по бусове"])
                self.assertEqual(saved["Статистики по бусове"].max_column, 11 + int(multiple_courses))
                self.assertEqual("Курс" in stats, multiple_courses)
                self.assertEqual(stats["Разстояние (км)"].value, routes[0].total_distance_km)
                self.assertEqual(stats["Капацитет използване (%)"].value, .5)
                self.assertEqual(stats["Капацитет използване (%)"].number_format, "0.0%")
                self.assertEqual(stats["Депо стартова точка"].value, "42.000000, 23.000000")
                self.assertEqual(stats["Крайна точка"].value, "42.000000, 23.000000")
                self.assertEqual(stats["Стартово време (чч:мм)"].value, "08:00")

    def test_excel_and_csv_exports_omit_center_distances_without_shifting_data(self):
        for multiple_courses in (False, True):
            with self.subTest(multiple_courses=multiple_courses), tempfile.TemporaryDirectory() as directory:
                self.exporter.config.excel_output_dir = directory
                self.exporter.config.csv_output_file = str(Path(directory) / "routes.csv")
                routes = [route([customer()])]
                if multiple_courses:
                    routes.append(route([customer("second")], trip_number=2, planned_start_minutes=600))
                solution = SimpleNamespace(routes=routes, total_distance_km=sum(r.total_distance_km for r in routes))
                path = self.exporter.export_all_to_single_excel(solution, [customer("unserved", coordinates=None)])
                saved = load_workbook(path)
                try:
                    self.assertEqual(len(saved.worksheets), 4)
                    for ws in saved.worksheets:
                        self.assertFalse(any("центъра" in str(cell.value).lower()
                                             for row in ws for cell in row))
                    self.assertEqual(self.by_label(saved["Обобщение"])["Общо разстояние (км)"],
                                     solution.total_distance_km)
                finally:
                    saved.close()
                path = self.exporter.export_routes_csv(solution)
                with open(path, encoding="utf-8-sig", newline="") as handle:
                    rows = list(csv.reader(handle))
                headers = rows[0]
                self.assertFalse(any("центъра" in header.lower() for header in headers))
                self.assertEqual(len(rows) - 1, len(routes))
                for row_values, current_route in zip(rows[1:], routes):
                    self.assertEqual(len(row_values), len(headers))
                    cells = dict(zip(headers, row_values))
                    self.assertEqual(cells["GPS координати"], current_route.customers[0].original_gps_data)
                    self.assertEqual(cells["Депо стартова точка"], "42.000000, 23.000000")
                    self.assertEqual(cells["Крайна точка"], "42.000000, 23.000000")
                    self.assertEqual(float(cells["Стартово време (мин)"]), current_route.planned_start_minutes)
                    expected = self.exporter._build_route_schedule_entries(current_route)[0]["distance_from_previous"]
                    self.assertEqual(float(cells["Разстояние от предишен (км)"]), round(expected, 2))
                    if multiple_courses:
                        self.assertEqual(int(cells["Курс"]), current_route.trip_number)

    def test_fallback_obeys_client_service_override_and_time_window_toggle(self):
        clients = [customer("zero", service_time_minutes=0, time_windows=[(600, 660)]),
                   customer("custom", service_time_minutes=7.5), customer("default")]
        for cls in (ExcelExporter, InteractiveMapGenerator):
            with self.subTest(exporter=cls.__name__):
                exporter = cls.__new__(cls)
                exporter.config = OutputConfig()
                with patch.object(exporter, "_calculate_time_between_points", return_value=0):
                    entries = exporter._build_route_schedule_entries(route(clients))
                    self.assertEqual([e["service_time_minutes"] for e in entries], [0, 7.5, 15])
                    self.assertEqual([e["wait_minutes"] for e in entries], [0, 0, 0])
                    self.assertEqual(entries[-1]["cumulative_time"], 22.5)
                    self.config.cvrp.enable_customer_time_windows = True
                    entries = exporter._build_route_schedule_entries(route(clients))
                    self.assertEqual(entries[0]["wait_minutes"], 120)
                    self.config.cvrp.enable_customer_time_windows = False

    def test_canonical_schedule_is_preserved_and_hhmm_is_rounded_consistently(self):
        client = customer(service_time_minutes=99)
        current_route = route([client], schedule_entries=[dict(
            customer=client, index=1, previous_stop_name="Депо", distance_from_previous=5.12345,
            cumulative_distance=5.12345, travel_time_minutes=10, service_time_minutes=7,
            wait_minutes=3, total_time_for_step=20.8, cumulative_time=20.8,
            start_time_minutes=480, arrival_time_minutes=493.8, total_time_with_start=500.8,
            time_window_text="", time_window_status="")])
        wb = self.workbook()
        self.exporter._create_routes_sheet(wb, SimpleNamespace(routes=[current_route]))
        cells = self.by_header(wb["Маршрути"])
        self.assertEqual(cells["Начало обслужване (ETA)"].value, "08:14")
        self.assertEqual(cells["Отпътуване от клиента (чч:мм)"].value, "08:21")
        self.assertAlmostEqual(cells["Пътуване + обслужване + чакане (ч)"].value, 20.8 / 60)
        self.assertEqual(cells["Разстояние от предишен (км)"].value, 5.12345)


if __name__ == "__main__":
    unittest.main()
